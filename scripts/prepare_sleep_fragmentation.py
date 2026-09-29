"""Convert the Nguyen, Fields & Ashley IL-6 workbook into pipeline inputs.

Source: doi:10.5061/dryad.tdz08kq3p (CC0), file ``Male_NSF_ASF_IL-6_ELISA_result.xlsx``,
kept unmodified in ``data/raw/sleep_fragmentation_il6/``.

The workbook is the authors' analysis spreadsheet, not a reader export. The
four plate grids are literal numbers, but the sample-to-well map and the
dilution factors exist only inside Excel formulas. For example, the sample
``M_ASF2h-1`` has an OD cell ``=AVERAGE(G30:H30)`` and a concentration cell
``=2*(K38-0.0097)/0.0144``. This script reads those formulas instead of
hand-typing a plate map, so every assignment can be traced to a cell:

* wells: the two cells named in each sample's ``AVERAGE(...)`` formula;
* dilution: a leading ``k*(`` multiplier in the sample's OD-blank or
  concentration formula (2× or 3×);
* standards: rows A-G of columns 1-2, with nominal levels from the column
  next to each grid;
* blank: the 0 pg/mL standard in H1/H2. The authors subtracted H9/H10
  (E9/E10 on plate 4). Those are undocumented unused wells. A zero standard
  is the correct blank for a sandwich ELISA, because it has been through
  every step except analyte.

Outputs go to ``data/processed/sleep_fragmentation_il6/``: grid-format plate
CSVs, layout CSVs, a ``manifest.csv``, per-sample metadata, and the authors'
own concentrations for comparison.

Usage::

    python scripts/prepare_sleep_fragmentation.py
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import column_index_from_string

REPO = Path(__file__).resolve().parents[1]
SOURCE = REPO / "data/raw/sleep_fragmentation_il6/Male_NSF_ASF_IL-6_ELISA_result.xlsx"
OUT = REPO / "data/processed/sleep_fragmentation_il6"

ROWS = "ABCDEFGH"
HEADER_ROWS = (1, 27, 58, 85)  # Excel rows holding "1 .. 12" for plates 1-4
FIRST_COL = column_index_from_string("C")  # Excel column C = plate column 1
SAMPLE_RE = re.compile(r"^M_(NSF|ASF)(\d+)h-(\d+)$")
AVERAGE_RE = re.compile(r"^=AVERAGE\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$")
MULT_RE = re.compile(r"^=(\d+(?:\.\d+)?)\*\(")


def excel_to_well(col_letters: str, row: int) -> tuple[int, str]:
    """Map an Excel cell inside a plate grid to ``(plate_number, well)``."""
    for plate, header in enumerate(HEADER_ROWS, start=1):
        if header < row <= header + 8:
            r = ROWS[row - header - 1]
            c = column_index_from_string(col_letters) - FIRST_COL + 1
            if not 1 <= c <= 12:
                break
            return plate, f"{r}{c}"
    raise ValueError(f"cell {col_letters}{row} is not inside a plate grid")


def main() -> None:
    wb_f = openpyxl.load_workbook(SOURCE, data_only=False)
    ws = wb_f["IL-6 Raw and SD"]

    # 1. Plate grids (literal OD values), written as reader-style exports.
    (OUT / "plates").mkdir(parents=True, exist_ok=True)
    (OUT / "layouts").mkdir(parents=True, exist_ok=True)
    grids: dict[int, list[list[float]]] = {}
    for plate, header in enumerate(HEADER_ROWS, start=1):
        assert [ws.cell(header, FIRST_COL + i).value for i in range(12)] == list(range(1, 13))
        grids[plate] = [
            [ws.cell(header + 1 + r, FIRST_COL + c).value for c in range(12)] for r in range(8)
        ]
        lines = [
            f"Source,doi:10.5061/dryad.tdz08kq3p (CC0) plate {plate}",
            "Analyte,mouse IL-6 (BioLegend ELISA MAX Deluxe 431304)",
            "",
            "," + ",".join(str(c) for c in range(1, 13)),
        ]
        lines += [f"{ROWS[r]}," + ",".join(f"{v}" for v in grids[plate][r]) for r in range(8)]
        (OUT / "plates" / f"plate{plate}.csv").write_text("\n".join(lines) + "\n")

    # 2. Standards: rows A-G of columns 1-2 on every plate. The nominal levels are the
    #    same on all plates (labels "S500" ... "S7.8" beside plate 1's grid, column Q).
    levels = [
        float(str(ws.cell(HEADER_ROWS[0] + 1 + r, column_index_from_string("Q")).value).lstrip("S"))
        for r in range(7)
    ]
    layout_rows: dict[int, list[dict]] = {p: [] for p in grids}
    for plate in grids:
        for r, level in enumerate(levels):
            for col in (1, 2):
                layout_rows[plate].append(
                    {
                        "well": f"{ROWS[r]}{col}",
                        "type": "standard",
                        "sample_id": f"STD{r + 1}",
                        "concentration": level,
                        "dilution": 1,
                        "group": "",
                    }
                )
        for col in (1, 2):  # the 0 pg/mL standard is the blank
            layout_rows[plate].append(
                {
                    "well": f"H{col}",
                    "type": "blank",
                    "sample_id": "",
                    "concentration": None,
                    "dilution": 1,
                    "group": "",
                }
            )

    # 3. Samples: wells + dilution recovered from each sample's formulas.
    samples = []
    for row in ws.iter_rows():
        for cell in row:
            m = SAMPLE_RE.match(str(cell.value or ""))
            if not m:
                continue
            treatment, time_h, rep = m.group(1), int(m.group(2)), int(m.group(3))
            od_formula = str(ws.cell(cell.row, cell.column + 1).value)
            a = AVERAGE_RE.match(od_formula)
            if not a:
                raise ValueError(f"{cell.coordinate}: unexpected OD formula {od_formula!r}")
            (p1, w1), (p2, w2) = (
                excel_to_well(a.group(1), int(a.group(2))),
                excel_to_well(a.group(3), int(a.group(4))),
            )
            assert p1 == p2
            dilution = 1.0
            for off in (2, 3):
                f = str(ws.cell(cell.row, cell.column + off).value or "")
                mm = MULT_RE.match(f)
                if mm:
                    dilution = float(mm.group(1))
            group = f"{treatment} {time_h}h"
            samples.append(
                {
                    "sample_id": cell.value,
                    "treatment": treatment,
                    "time_h": time_h,
                    "replicate": rep,
                    "group": group,
                    "plate": p1,
                    "wells": f"{w1},{w2}",
                    "dilution": dilution,
                    "source_cell": cell.coordinate,
                }
            )
            for w in (w1, w2):
                layout_rows[p1].append(
                    {
                        "well": w,
                        "type": "sample",
                        "sample_id": cell.value,
                        "concentration": None,
                        "dilution": dilution,
                        "group": group,
                        "treatment": treatment,
                        "time_h": time_h,
                    }
                )

    sample_df = pd.DataFrame(samples).sort_values(["plate", "treatment", "time_h", "replicate"])
    if len(sample_df) != 110 or sample_df["sample_id"].duplicated().any():
        raise RuntimeError(f"expected 110 unique samples, found {len(sample_df)}")
    sample_df.to_csv(OUT / "sample_info.csv", index=False)

    manifest = []
    for plate, rows in layout_rows.items():
        lay = pd.DataFrame(rows)
        if lay["well"].duplicated().any():
            raise RuntimeError(f"plate {plate}: a well is assigned twice")
        lay.to_csv(OUT / "layouts" / f"plate{plate}_layout.csv", index=False)
        manifest.append(
            {
                "plate_id": f"plate{plate}",
                "analyte": "IL-6",
                "plate_file": f"plates/plate{plate}.csv",
                "layout_file": f"layouts/plate{plate}_layout.csv",
                "units": "pg/mL",
            }
        )
    pd.DataFrame(manifest).to_csv(OUT / "manifest.csv", index=False)

    # 4. The authors' own concentrations (linear fits), for comparison.
    wb_v = openpyxl.load_workbook(SOURCE, data_only=True)
    an = wb_v["IL-6 Analyzed and graph"]
    authors = [
        {"sample_id": an.cell(r, 1).value, "authors_conc": an.cell(r, 2).value}
        for r in range(2, an.max_row + 1)
        if isinstance(an.cell(r, 1).value, str) and SAMPLE_RE.match(an.cell(r, 1).value)
    ]
    pd.DataFrame(authors).to_csv(OUT / "authors_concentrations.csv", index=False)

    print(f"wrote 4 plates, {len(sample_df)} samples -> {OUT.relative_to(REPO)}")
    print(sample_df.groupby("dilution").size().rename("samples by dilution").to_string())


if __name__ == "__main__":
    main()
